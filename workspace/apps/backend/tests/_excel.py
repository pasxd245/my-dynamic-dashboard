"""Build small Excel workbooks in-memory for upload-flow tests."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook


def two_sheet_workbook() -> bytes:
    wb = Workbook()
    deals = wb.active
    deals.title = "Deals"
    deals.append(["deal_id", "amount", "closed_at"])
    deals.append(["d1", 100, "2024-01-15"])
    deals.append(["d2", 250, "2024-02-03"])

    contacts = wb.create_sheet("Contacts")
    contacts.append(["contact_id", "email"])
    contacts.append(["c1", "alice@example.com"])
    contacts.append(["c2", "bob@example.com"])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def malformed_workbook() -> bytes:
    """A bytes blob that is not actually an xlsx — used for 415 path."""
    return b"not really an xlsx"


def fm1_shaped_workbook() -> bytes:
    """R144 — the FM1 "Weekly - Report Call" shape: an agent column + a TEXT
    `dd-MM-yyyy HH:mm:ss` call timestamp (pandas reads `object` → the parser
    infers `string`, exactly like the real `Ngày gọi`), spanning two ISO
    weeks; plus a sheet with one unparseable cell for the 422 path."""
    wb = Workbook()
    calls = wb.active
    calls.title = "Calls"
    calls.append(["agent", "called_at", "call_date"])
    calls.append(["An", "29-06-2026 09:15:00", "29-06-2026"])  # Mon — ISO week starting 2026-06-29
    calls.append(["An", "01-07-2026 10:30:00", "01-07-2026"])  # Wed, same week
    calls.append(["Binh", "03-07-2026 17:45:00", "03-07-2026"])  # Fri, same week
    calls.append(["An", "06-07-2026 08:00:00", "06-07-2026"])  # Mon — the NEXT ISO week
    calls.append(["Binh", None, None])  # NULL passes through

    bad = wb.create_sheet("BadDates")
    bad.append(["agent", "called_at"])
    bad.append(["An", "29-06-2026 09:15:00"])
    bad.append(["Chi", "not a date"])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def mixed_type_workbook() -> bytes:
    """R143 — the FM2.25 shape: a phone column mixing numeric cells with
    leading-zero TEXT cells (pandas reads `object`; pre-R143 the parquet
    write died with ArrowInvalid), plus a clean sheet for atomicity tests."""
    wb = Workbook()
    calls = wb.active
    calls.title = "Calls"
    calls.append(["caller", "phone", "duration"])
    calls.append(["a", 903359280, 10])
    calls.append(["b", "0387353189", 20])  # leading-zero text cell
    calls.append(["c", "09-8157-2157", 30])  # non-numeric text keeps the column mixed
    calls.append(["d", None, 40])  # NULL passes through

    clean = wb.create_sheet("Clean")
    clean.append(["k", "v"])
    clean.append(["x", 1])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def blank_cell_workbook() -> bytes:
    """R175 — the CRM-export shape: a TEXT `dd-MM-yyyy HH:mm:ss` timestamp and a
    numeric column, each carrying a WHITESPACE-ONLY cell (the real export had one
    `" "` in 31k rows that aborted the whole commit), plus a string column whose
    blank text must survive verbatim."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    ws.append(["called_at", "amount", "note"])
    ws.append(["29-06-2026 09:15:00", "100", "a"])
    ws.append([" ", "  ", " "])
    ws.append(["01-07-2026 10:30:00", "\t", "b"])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def leading_zero_text_workbook() -> bytes:
    """R175 — a phone column stored as TEXT with a leading zero in every cell
    (the TNDS export shape). pandas' reader turns the whole column into int64."""
    wb = Workbook()
    ws = wb.active
    ws.title = "TNDS"
    ws.append(["phone", "contract"])
    for phone, contract in (("0968733440", "C1"), ("0915034727", "C2"), ("0378932776", "C3")):
        ws.append([phone, contract])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
