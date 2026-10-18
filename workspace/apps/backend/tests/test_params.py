"""R194 (item 3.7a) — typed query parameters: the engine.

Design `dashboards/dashboard-filters.md` criteria: 1 precedence (viewer >
dashboard > default > `param_unbound`); 2 `param_invalid` and the save checks;
3 two values → two answers, one `definitionVersion`, a query without parameters
unchanged; 4 a relative default in `backend.time_zone` across midnight; 7 totals
and export carry the values; 8 workflow sources run on defaults and stamp them.
Plus the contract-exit rulings: only USED parameters resolve / echo (F2); a date
parameter on a datetime column is a calendar day (F3); the dashboard level is
read by the server (F1).
"""

from __future__ import annotations

import io
import json
from datetime import datetime, timezone

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import db, params
from app._config import CONFIG
from app.main import app
from app.metadata import definition_version
from tests._conformance import validate_response

_CSV = (
    b"id,team,amount,day,at\n"
    b"1,North,10,2026-10-10,2026-10-10 09:00:00\n"
    b"2,North,20,2026-10-11,2026-10-11 18:30:00\n"
    b"3,South,5,2026-10-12,2026-10-12 08:00:00\n"
)
_EMPTY = {"q": None, "filters": [], "advanced": []}


def _dataset(client: TestClient) -> tuple[str, str, list[dict]]:
    ws = client.post("/workspaces", json={"name": "P"}).json()["id"]
    up = client.post("/uploads", data={"sourceFormat": "csv"}, files={"file": ("p.csv", _CSV, "text/csv")}).json()
    ds = client.post(
        f"/workspaces/{ws}/datasets/batch", json={"temp_id": up["temp_id"], "items": [{"name": "calls"}]}
    ).json()[0]
    return ws, ds["id"], ds["columns"]


def _col(cols: list[dict], name: str) -> int:
    return next(i for i, c in enumerate(cols) if c["name"] == name)


def _query(client: TestClient, ws: str, ds: str, definition: dict, name: str = "q"):
    return client.post(f"/workspaces/{ws}/queries", json={"name": name, "sourceId": ds, "definition": definition})


def _as_of_query(cols: list[dict], default=None, label: str | None = "As of") -> dict:
    p: dict = {"name": "as_of", "type": "date"}
    if label:
        p["label"] = label
    if default is not None:
        p["default"] = default
    return {
        **_EMPTY,
        "params": [p],
        "filters": [{"col": _col(cols, "day"), "dtype": "date", "op": "lte", "val": {"param": "as_of"}}],
    }


def _ids(rows_body: dict, cols: list[dict]) -> list[str]:
    return sorted(r[_col(cols, "id")] for r in rows_body["rows"])


# --- 1 · precedence ------------------------------------------------------------


@pytest.mark.unit
def test_precedence_viewer_dashboard_default_unbound() -> None:
    with TestClient(app) as client:
        ws, ds, cols = _dataset(client)
        q_default = _query(client, ws, ds, _as_of_query(cols, default="2026-10-10"), "with default").json()["id"]
        q_bare = _query(client, ws, ds, _as_of_query(cols), "no default").json()["id"]
        dsh = client.post(
            f"/workspaces/{ws}/dashboards",
            json={
                "name": "D",
                "slug": "d",
                "definition": {
                    "widgets": [
                        {
                            "id": "wdg_00000001",
                            "queryId": q_bare,
                            "title": "T",
                            "chartType": "table",
                            "agg": "count",
                            "span": 3,
                        }
                    ]
                },
            },
        ).json()["id"]
        # The saved control values arrive with 3.7b's `controls`; written here directly.
        with db.get_conn() as con:
            row = con.execute("SELECT definition_json FROM dashboards WHERE id = ?", (dsh,)).fetchone()
            d = json.loads(row["definition_json"])
            d["controls"] = [
                {"id": "ctl_00000001", "kind": "param", "param": "as_of", "targets": "all", "saved": "2026-10-11"}
            ]
            con.execute("UPDATE dashboards SET definition_json = ? WHERE id = ?", (json.dumps(d), dsh))
            con.commit()
        dash = f"dashboard={dsh}&widget=wdg_00000001"

        viewer = client.get(f"/queries/{q_bare}/rows?p.as_of=2026-10-12&{dash}")
        dashboard = client.get(f"/queries/{q_bare}/rows?{dash}")
        default = client.get(f"/queries/{q_default}/rows")
        unbound = client.get(f"/queries/{q_bare}/rows")
        agg_dash = client.post(
            f"/queries/{q_bare}/aggregate",
            json={
                "dimensions": [],
                "measures": [{"agg": "count"}],
                "dashboard": {"id": dsh, "widgetId": "wdg_00000001"},
            },
        )

    assert viewer.status_code == 200, viewer.text
    assert viewer.json()["resolvedParams"] == {"as_of": {"value": "2026-10-12", "source": "viewer"}}
    assert _ids(viewer.json(), cols) == ["1", "2", "3"]
    assert dashboard.json()["resolvedParams"] == {"as_of": {"value": "2026-10-11", "source": "dashboard"}}
    assert _ids(dashboard.json(), cols) == ["1", "2"]
    assert default.json()["resolvedParams"] == {"as_of": {"value": "2026-10-10", "source": "default"}}
    assert _ids(default.json(), cols) == ["1"]
    assert unbound.status_code == 422
    assert unbound.json() == {"code": "param_unbound", "param": "as_of", "label": "As of"}
    validate_response("queries/rows-get.contract.yaml", 200, viewer.json())
    validate_response("queries/rows-get.contract.yaml", 422, unbound.json())
    assert agg_dash.json()["rows"] == [["2"]] and agg_dash.json()["resolvedParams"]["as_of"]["source"] == "dashboard"
    validate_response("queries/aggregate.contract.yaml", 200, agg_dash.json())


@pytest.mark.unit
def test_an_unknown_dashboard_or_widget_is_refused() -> None:
    with TestClient(app) as client:
        ws, ds, cols = _dataset(client)
        q = _query(client, ws, ds, _as_of_query(cols, default="2026-10-10")).json()["id"]
        unknown = client.get(f"/queries/{q}/rows?dashboard=dsh_00000000&widget=wdg_00000001")
        half = client.get(f"/queries/{q}/rows?dashboard=dsh_00000000")

    assert unknown.status_code == 422 and "unknown_dashboard" in unknown.text
    assert half.status_code == 422


# --- 2 · invalid values and save checks ----------------------------------------------


@pytest.mark.unit
def test_a_wrong_typed_value_is_param_invalid() -> None:
    with TestClient(app) as client:
        ws, ds, cols = _dataset(client)
        q = _query(client, ws, ds, _as_of_query(cols, default="2026-10-10")).json()["id"]
        bad_get = client.get(f"/queries/{q}/rows?p.as_of=last-week")
        bad_post = client.post(f"/queries/{q}/totals", json={"columns": ["amount"], "params": {"as_of": 5}})
        iso_week = client.get(f"/queries/{q}/rows?p.as_of=2026-W41-1")

    body = {"code": "param_invalid", "param": "as_of", "label": "As of", "expected": "date", "source": "viewer"}
    assert bad_get.status_code == 422 and bad_get.json() == body
    assert bad_post.status_code == 422 and bad_post.json() == body
    assert iso_week.status_code == 422
    validate_response("queries/totals.contract.yaml", 422, bad_post.json())


@pytest.mark.unit
def test_save_checks_name_the_field() -> None:
    with TestClient(app) as client:
        ws, ds, cols = _dataset(client)
        day, team, amount, at = (_col(cols, n) for n in ("day", "team", "amount", "at"))
        p_date = [{"name": "as_of", "type": "date"}]

        def save(definition: dict):
            return _query(client, ws, ds, {**_EMPTY, **definition}, name=f"q{id(definition)}")

        undeclared = save({"filters": [{"col": day, "dtype": "date", "op": "lte", "val": {"param": "nope"}}]})
        wrong_type = save(
            {"params": p_date, "filters": [{"col": amount, "dtype": "integer", "op": "gt", "val": {"param": "as_of"}}]}
        )
        no_operand = save(
            {"params": p_date, "filters": [{"col": day, "dtype": "date", "op": "is_null", "val": {"param": "as_of"}}]}
        )
        with_ref_col = save(
            {
                "params": p_date,
                "filters": [{"col": day, "dtype": "date", "op": "lte", "val": {"param": "as_of"}, "refCol": day}],
            }
        )
        in_then = save(
            {
                "params": [{"name": "t", "type": "string"}],
                "steps": [
                    {
                        "kind": "case",
                        "name": "c",
                        "branches": [
                            {
                                "when": [{"col": "team", "op": "equals", "val": "North"}],
                                "then": {"kind": "const", "value": {"param": "t"}},
                            }
                        ],
                    }
                ],
            }
        )
        in_derive = save(
            {
                "params": [{"name": "n", "type": "number"}],
                "steps": [{"kind": "derive", "name": "x", "left": "amount", "op": "add", "right": {"param": "n"}}],
            }
        )
        day_match = save(
            {"params": p_date, "filters": [{"col": at, "dtype": "datetime", "op": "equals", "val": {"param": "as_of"}}]}
        )
        dup = save({"params": [*p_date, *p_date]})
        bad_default = save({"params": [{"name": "n", "type": "number", "default": "ten"}]})
        rel_on_number = save(
            {"params": [{"name": "n", "type": "number", "default": {"kind": "today", "offsetDays": 0}}]}
        )
        step_type = save(
            {
                "params": p_date,
                "steps": [
                    {"kind": "filter", "predicates": [{"col": "team", "op": "equals", "val": {"param": "as_of"}}]}
                ],
            }
        )
        too_many = save({"params": [{"name": f"p{i}", "type": "string"} for i in range(9)]})
        unused_ok = save({"params": p_date})
        ok_step = save(
            {
                "params": [{"name": "who", "type": "string", "default": "North"}],
                "steps": [{"kind": "filter", "predicates": [{"col": "team", "op": "equals", "val": {"param": "who"}}]}],
            }
        )
        _ = team

    def msg(r) -> str:
        assert r.status_code == 422, r.text
        return json.dumps(r.json())

    assert "param_undeclared" in msg(undeclared) and '"filters", 0, "val"' in msg(undeclared)
    assert "param_type" in msg(wrong_type)
    assert "param_operand" in msg(no_operand)
    assert "param_ref_col" in msg(with_ref_col)
    msg(in_then)  # a reference is not a CaseValue (D2: no)
    msg(in_derive)  # nor a derive constant
    assert "param_day_match" in msg(day_match)
    assert "param_duplicate" in msg(dup)
    assert "param_default" in msg(bad_default)
    assert "param_default" in msg(rel_on_number)
    assert "param_type" in msg(step_type) and '"steps", 0, "predicates", 0, "val"' in msg(step_type)
    msg(too_many)
    assert unused_ok.status_code == 201, unused_ok.text
    assert ok_step.status_code == 201, ok_step.text


# --- 3 · values never touch the definition ----------------------------------------------


@pytest.mark.unit
def test_two_values_two_answers_one_definition_version() -> None:
    with TestClient(app) as client:
        ws, ds, cols = _dataset(client)
        created = _query(client, ws, ds, _as_of_query(cols, default="2026-10-10"))
        q = created.json()["id"]
        a = client.get(f"/queries/{q}/rows?p.as_of=2026-10-10")
        b = client.get(f"/queries/{q}/rows?p.as_of=2026-10-12")
        after = client.get(f"/queries/{q}").json()
        with db.get_conn() as con:
            stored = json.loads(con.execute("SELECT definition_json FROM queries WHERE id = ?", (q,)).fetchone()[0])

    assert a.json()["total"] == 1 and b.json()["total"] == 3
    assert after["definition"] == created.json()["definition"]
    assert stored["filters"][0]["val"] == {"param": "as_of"}
    assert definition_version(stored) == definition_version(json.loads(json.dumps(after["definition"])))
    validate_response("queries/detail-get.contract.yaml", 200, after)


@pytest.mark.unit
def test_a_query_without_parameters_answers_as_before() -> None:
    with TestClient(app) as client:
        ws, ds, cols = _dataset(client)
        q = _query(
            client,
            ws,
            ds,
            {**_EMPTY, "filters": [{"col": _col(cols, "team"), "dtype": "string", "op": "equals", "val": "North"}]},
        ).json()["id"]
        detail = client.get(f"/queries/{q}").json()
        rows = client.get(f"/queries/{q}/rows?p.as_of=2026-10-10").json()
        totals = client.post(f"/queries/{q}/totals", json={"columns": ["amount"], "params": {"x": "y"}}).json()
        csv = client.post(f"/queries/{q}/export", json={"format": "csv"})
        with db.get_conn() as con:
            stored = json.loads(con.execute("SELECT definition_json FROM queries WHERE id = ?", (q,)).fetchone()[0])

    assert "params" not in detail["definition"]
    assert stored["params"] is None  # stored as null: the hash drops it
    assert set(rows) == {"rows", "page", "pageSize", "total"}
    assert totals == {"sums": {"amount": "30"}, "rowCount": 2}
    assert "x-mdd-params" not in {k.lower() for k in csv.headers}


@pytest.mark.unit
def test_only_used_parameters_resolve_and_echo() -> None:
    with TestClient(app) as client:
        ws, ds, cols = _dataset(client)
        d = _as_of_query(cols, default="2026-10-11")
        d["params"].append({"name": "unused", "type": "string"})  # no default, never referenced
        q = _query(client, ws, ds, d).json()["id"]
        r = client.get(f"/queries/{q}/rows?p.unused=hello")

    assert r.status_code == 200
    assert r.json()["resolvedParams"] == {"as_of": {"value": "2026-10-11", "source": "default"}}


# --- 4 · today in the configured zone ---------------------------------------------------


@pytest.mark.unit
def test_relative_default_resolves_in_the_configured_zone(monkeypatch: pytest.MonkeyPatch) -> None:
    # 23:30 UTC on 11 Oct is already 12 Oct in Ho Chi Minh City (UTC+7).
    monkeypatch.setattr(params, "_clock", lambda: datetime(2026, 10, 11, 23, 30, tzinfo=timezone.utc))
    with TestClient(app) as client:
        ws, ds, cols = _dataset(client)
        q = _query(client, ws, ds, _as_of_query(cols, default={"kind": "today", "offsetDays": -1})).json()["id"]
        utc = client.get(f"/queries/{q}/rows").json()
        monkeypatch.setattr(CONFIG.settings.backend, "time_zone", "Asia/Ho_Chi_Minh")
        hcm = client.get(f"/queries/{q}/rows").json()

    assert utc["resolvedParams"]["as_of"] == {"value": "2026-10-10", "source": "default"}
    assert utc["total"] == 1
    assert hcm["resolvedParams"]["as_of"] == {"value": "2026-10-11", "source": "default"}
    assert hcm["total"] == 2


# --- F3 · a date parameter on a datetime column is a calendar day ---------------------------


@pytest.mark.unit
def test_a_date_parameter_on_a_datetime_column_covers_the_whole_day() -> None:
    with TestClient(app) as client:
        ws, ds, cols = _dataset(client)
        assert cols[_col(cols, "at")]["dtype"] == "datetime"
        at = _col(cols, "at")
        defs = {
            "lte": {"op": "lte", "val": {"param": "d"}},
            "gte": {"op": "gte", "val": {"param": "d"}},
            "after": {"op": "after", "val": {"param": "d"}},
            "before": {"op": "before", "val": {"param": "d"}},
            "between": {"op": "between", "min": {"param": "d"}, "max": {"param": "d"}},
            "between_lit": {"op": "between", "min": "2026-10-10 00:00:00", "max": {"param": "d"}},
        }
        out = {}
        for key, atom in defs.items():
            definition = {
                **_EMPTY,
                "params": [{"name": "d", "type": "date"}],
                "filters": [{"col": at, "dtype": "datetime", **atom}],
            }
            created = _query(client, ws, ds, definition, name=key)
            assert created.status_code == 201, (key, created.text)
            out[key] = _ids(client.get(f"/queries/{created.json()['id']}/rows?p.d=2026-10-11").json(), cols)

    assert out["lte"] == ["1", "2"]  # 18:30 on the 11th is in
    assert out["gte"] == ["2", "3"]
    assert out["after"] == ["3"]  # after the whole day
    assert out["before"] == ["1"]
    assert out["between"] == ["2"]
    assert out["between_lit"] == ["1", "2"]


# --- steps, case WHEN, preview ------------------------------------------------------------


@pytest.mark.unit
def test_parameters_in_filter_steps_case_when_and_preview() -> None:
    with TestClient(app) as client:
        ws, ds, cols = _dataset(client)
        definition = {
            **_EMPTY,
            "params": [{"name": "who", "type": "string", "default": "North"}, {"name": "min", "type": "number"}],
            "steps": [
                {"kind": "filter", "predicates": [{"col": "amount", "op": "gte", "val": {"param": "min"}}]},
                {
                    "kind": "case",
                    "name": "mine",
                    "branches": [
                        {
                            "when": [{"col": "team", "op": "in", "vals": [{"param": "who"}, "Nobody"]}],
                            "then": {"kind": "const", "value": "yes"},
                        }
                    ],
                    "otherwise": {"kind": "const", "value": "no"},
                },
            ],
        }
        created = _query(client, ws, ds, definition)
        q = created.json()["id"]
        rows = client.get(f"/queries/{q}/rows?p.min=10").json()
        detail = client.get(f"/queries/{q}").json()
        preview = client.post(
            f"/workspaces/{ws}/queries/preview",
            json={"sourceId": ds, "definition": definition, "params": {"min": "5", "who": "South"}},
        )
        preview_unbound = client.post(
            f"/workspaces/{ws}/queries/preview", json={"sourceId": ds, "definition": definition}
        )

    assert created.status_code == 201, created.text
    assert rows["total"] == 2
    assert {r[-1] for r in rows["rows"]} == {"yes"}
    assert rows["resolvedParams"] == {
        "min": {"value": 10, "source": "viewer"},
        "who": {"value": "North", "source": "default"},
    }
    assert [c["name"] for c in detail["resolvedColumns"]][-1] == "mine"
    assert preview.status_code == 200, preview.text
    assert preview.json()["total"] == 3 and preview.json()["resolvedParams"]["who"]["value"] == "South"
    validate_response("queries/preview.contract.yaml", 200, preview.json())
    assert preview_unbound.status_code == 422 and preview_unbound.json()["param"] == "min"


# --- 7 · what leaves the app ---------------------------------------------------------------


@pytest.mark.unit
def test_totals_and_export_carry_the_values() -> None:
    with TestClient(app) as client:
        ws, ds, cols = _dataset(client)
        q = _query(client, ws, ds, _as_of_query(cols, default="2026-10-10")).json()["id"]
        body = {"params": {"as_of": "2026-10-11"}}
        totals = client.post(f"/queries/{q}/totals", json={"columns": ["amount"], **body})
        xlsx = client.post(f"/queries/{q}/export", json={"format": "xlsx", **body})
        csv = client.post(f"/queries/{q}/export", json={"format": "csv", **body})

    assert totals.json() == {
        "sums": {"amount": "30"},
        "rowCount": 2,
        "resolvedParams": {"as_of": {"value": "2026-10-11", "source": "viewer"}},
    }
    validate_response("queries/totals.contract.yaml", 200, totals.json())
    header = {"as_of": {"value": "2026-10-11", "source": "viewer"}}
    assert json.loads(csv.headers["x-mdd-params"]) == header
    assert json.loads(xlsx.headers["x-mdd-params"]) == header
    fs = load_workbook(io.BytesIO(xlsx.content))["Filters"]
    cells = [[c.value for c in r] for r in fs.iter_rows()]
    i = next(k for k, r in enumerate(cells) if r[0] == "Parameters")
    assert cells[i + 1][:3] == ["name", "value", "source"]
    assert cells[i + 2][:3] == ["as_of", "2026-10-11", "viewer"]


@pytest.mark.unit
def test_the_export_header_is_ascii_json() -> None:
    with TestClient(app) as client:
        ws, ds, cols = _dataset(client)
        d = {
            **_EMPTY,
            "params": [{"name": "who", "type": "string"}],
            "filters": [{"col": _col(cols, "team"), "dtype": "string", "op": "ne", "val": {"param": "who"}}],
        }
        q = _query(client, ws, ds, d).json()["id"]
        csv = client.post(f"/queries/{q}/export", json={"format": "csv", "params": {"who": "Đông Nam"}})

    raw = csv.headers["x-mdd-params"]
    assert raw.isascii()
    assert json.loads(raw)["who"]["value"] == "Đông Nam"


# --- 8 · workflow runs ------------------------------------------------------------------------


@pytest.mark.unit
def test_a_workflow_runs_its_sources_on_defaults_and_stamps_them() -> None:
    with TestClient(app) as client:
        ws, ds, cols = _dataset(client)
        q = _query(client, ws, ds, _as_of_query(cols, default="2026-10-11")).json()["id"]
        wf = client.post(
            f"/workspaces/{ws}/workflows", json={"name": "wf", "definition": {"sources": [q], "steps": []}}
        )
        run = client.post(f"/workflows/{wf.json()['id']}/run")
        rows = client.get(f"/workflows/{wf.json()['id']}/rows").json()
        step_ref = client.post(
            f"/workspaces/{ws}/workflows",
            json={
                "name": "wf3",
                "definition": {
                    "sources": [q],
                    "steps": [
                        {"kind": "filter", "predicates": [{"col": "team", "op": "equals", "val": {"param": "as_of"}}]}
                    ],
                },
            },
        )
        # The default is removed afterwards: the next run refuses, naming both.
        d = _as_of_query(cols)
        client.put(f"/queries/{q}", json={"definition": d})
        refused = client.post(f"/workflows/{wf.json()['id']}/run")
        bare = _query(client, ws, ds, _as_of_query(cols), "bare").json()["id"]
        save_bare = client.post(
            f"/workspaces/{ws}/workflows", json={"name": "wf2", "definition": {"sources": [bare], "steps": []}}
        )

    assert wf.status_code == 201, wf.text
    assert run.status_code == 200, run.text
    assert run.json()["run"]["sources"][0]["params"] == {"as_of": "2026-10-11"}
    validate_response("workflows/run-post.contract.yaml", 200, run.json())
    assert rows["total"] == 2
    assert refused.status_code == 422
    assert refused.json() == {"code": "param_unbound", "param": "as_of", "label": "As of", "sourceId": q}
    validate_response("workflows/run-post.contract.yaml", 422, refused.json())
    assert save_bare.status_code == 422 and "as_of" in save_bare.text
    assert step_ref.status_code == 422 and "param_in_workflow" in step_ref.text


# --- review round ------------------------------------------------------------------------


@pytest.mark.unit
def test_numbers_on_an_integer_column_must_be_whole_and_in_range() -> None:
    with TestClient(app) as client:
        ws, ds, cols = _dataset(client)
        assert cols[_col(cols, "amount")]["dtype"] == "integer"
        n_atom = {"col": _col(cols, "amount"), "dtype": "integer", "op": "gte", "val": {"param": "n"}}
        q = _query(client, ws, ds, {**_EMPTY, "params": [{"name": "n", "type": "number"}], "filters": [n_atom]}).json()[
            "id"
        ]
        whole_float = client.get(f"/queries/{q}/rows?p.n=1e1")
        half = client.get(f"/queries/{q}/rows?p.n=5.5")
        huge = client.get(f"/queries/{q}/rows?p.n=" + "9" * 40)
        underscore = client.get(f"/queries/{q}/rows?p.n=1_0")
        boolean = client.post(f"/queries/{q}/totals", json={"columns": ["amount"], "params": {"n": True}})
        bad_default = _query(
            client,
            ws,
            ds,
            {**_EMPTY, "params": [{"name": "n", "type": "number", "default": 2.5}], "filters": [n_atom]},
            "bad default",
        )

    assert whole_float.status_code == 200 and whole_float.json()["total"] == 2
    for r in (half, huge, underscore):
        assert r.status_code == 422 and r.json()["code"] == "param_invalid", r.text
    assert boolean.status_code == 422  # `true` is not 1
    assert bad_default.status_code == 422 and "param_default" in bad_default.text


@pytest.mark.unit
def test_a_reversed_range_is_the_values_fault() -> None:
    with TestClient(app) as client:
        ws, ds, cols = _dataset(client)
        d = {
            **_EMPTY,
            "params": [{"name": "a", "type": "date", "label": "From"}, {"name": "b", "type": "date", "label": "To"}],
            "filters": [
                {
                    "col": _col(cols, "day"),
                    "dtype": "date",
                    "op": "between",
                    "min": {"param": "a"},
                    "max": {"param": "b"},
                }
            ],
        }
        q = _query(client, ws, ds, d).json()["id"]
        reversed_ = client.get(f"/queries/{q}/rows?p.a=2026-10-12&p.b=2026-10-10")
        both_missing = client.get(f"/queries/{q}/rows")
        cleared = client.get(f"/queries/{q}/rows?p.a=&p.b=2026-10-10")

    assert reversed_.status_code == 422
    assert reversed_.json() == {
        "code": "param_invalid",
        "param": "a",
        "label": "From",
        "expected": "date",
        "source": "viewer",
        "reason": "range_reversed",
    }
    validate_response("queries/rows-get.contract.yaml", 422, reversed_.json())
    assert both_missing.json() == {
        "code": "param_unbound",
        "param": "a",
        "label": "From",
        "missing": [{"param": "a", "label": "From"}, {"param": "b", "label": "To"}],
    }
    validate_response("queries/rows-get.contract.yaml", 422, both_missing.json())
    assert cleared.status_code == 422 and cleared.json()["param"] == "a"  # empty = no value


@pytest.mark.unit
def test_the_dashboard_level_needs_a_widget_showing_this_query_and_beats_the_default() -> None:
    with TestClient(app) as client:
        ws, ds, cols = _dataset(client)
        q = _query(client, ws, ds, _as_of_query(cols, default="2026-10-10")).json()["id"]
        other = _query(client, ws, ds, _as_of_query(cols, default="2026-10-10"), "other").json()["id"]
        dsh = client.post(
            f"/workspaces/{ws}/dashboards",
            json={
                "name": "D",
                "slug": "d",
                "definition": {
                    "widgets": [
                        {
                            "id": "wdg_00000001",
                            "queryId": q,
                            "title": "T",
                            "chartType": "table",
                            "agg": "count",
                            "span": 3,
                        }
                    ]
                },
            },
        ).json()["id"]
        with db.get_conn() as con:
            d = json.loads(con.execute("SELECT definition_json FROM dashboards WHERE id = ?", (dsh,)).fetchone()[0])
            d["controls"] = [
                {
                    "id": "ctl_00000001",
                    "kind": "param",
                    "param": "as_of",
                    "targets": ["wdg_00000001"],
                    "saved": {"kind": "today", "offsetDays": 0},
                }
            ]
            con.execute("UPDATE dashboards SET definition_json = ? WHERE id = ?", (json.dumps(d), dsh))
            con.commit()
        mine = client.get(f"/queries/{q}/rows?dashboard={dsh}&widget=wdg_00000001")
        borrowed = client.get(f"/queries/{other}/rows?dashboard={dsh}&widget=wdg_00000001")

    assert mine.json()["resolvedParams"]["as_of"]["source"] == "dashboard"  # beats the default
    assert borrowed.status_code == 422 and "unknown_widget" in borrowed.text


@pytest.mark.unit
def test_a_published_dataset_carries_the_values_in_its_lineage() -> None:
    with TestClient(app) as client:
        ws, ds, cols = _dataset(client)
        q = _query(client, ws, ds, _as_of_query(cols, default="2026-10-11")).json()["id"]
        wf = client.post(
            f"/workspaces/{ws}/workflows", json={"name": "wf", "definition": {"sources": [q], "steps": []}}
        )
        client.post(f"/workflows/{wf.json()['id']}/run")
        pub = client.post(f"/workflows/{wf.json()['id']}/publish")

    assert pub.status_code == 201, pub.text
    assert pub.json()["lineage"]["sources"][0]["params"] == {"as_of": "2026-10-11"}


@pytest.mark.unit
def test_an_export_title_in_vietnamese_downloads() -> None:
    with TestClient(app) as client:
        ws, ds, cols = _dataset(client)
        q = _query(client, ws, ds, _EMPTY).json()["id"]
        r = client.post(f"/queries/{q}/export", json={"format": "csv", "title": "Đơn hàng"})

    assert r.status_code == 200
    disp = r.headers["content-disposition"]
    assert disp.startswith('attachment; filename="') and "filename*=UTF-8''" in disp
