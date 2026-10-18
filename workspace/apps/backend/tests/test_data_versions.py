"""R196 (item 3.2c) — every read says which data it read.

`dataVersions: {<datasetId>: n}` on `/rows`, `/aggregate`, `/totals` and the
preview — the versions of exactly the datasets the read read (a join: both) —
and in the export (a `Data versions` block on the Filters sheet, plus the
`X-MDD-Data-Versions` header). And `POST /datasets/{id}/export`: a dataset's
(or a saved copy's) rows as a file, with its version.
"""

from __future__ import annotations

import io
import json
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from tests._conformance import validate_response
from tests.test_joins import _create_join, _declare_id_join, _joined_def, _seed

_FIXTURES = Path(__file__).parent / "fixtures"
_EMPTY = {"q": None, "filters": [], "advanced": []}


def _refresh(client: TestClient, ws: str, ds: str) -> None:
    up = client.post(
        "/uploads",
        data={"sourceFormat": "csv"},
        files={"file": ("sample.csv", (_FIXTURES / "sample.csv").read_bytes(), "text/csv")},
    ).json()
    r = client.post(
        f"/workspaces/{ws}/datasets/batch",
        json={"temp_id": up["temp_id"], "items": [{"name": "x", "target_dataset_id": ds}]},
    )
    assert r.status_code == 201, r.text


def _meta(wb) -> dict:
    """The Filters sheet's leading key/value rows (up to the first blank row)."""
    out: dict = {}
    for r in wb["Filters"].iter_rows():
        if r[0].value is None:
            break
        out[r[0].value] = r[1].value
    return out


@pytest.mark.unit
def test_each_read_reports_the_versions_of_exactly_the_datasets_it_read() -> None:
    with TestClient(app) as client:
        ws, deals, accounts = _seed(client)
        single = client.post(
            f"/workspaces/{ws}/queries", json={"name": "deals", "sourceId": deals, "definition": _EMPTY}
        ).json()["id"]
        rel = _declare_id_join(client, ws, deals, accounts)
        joined = _create_join(client, ws, deals, rel).json()["id"]
        _refresh(client, ws, accounts)  # accounts → v2

        rows = client.get(f"/queries/{single}/rows").json()
        jrows = client.get(f"/queries/{joined}/rows").json()
        agg = client.post(
            f"/queries/{joined}/aggregate", json={"dimensions": [], "measures": [{"agg": "count"}]}
        ).json()
        jcols = client.get(f"/queries/{joined}").json()["resolvedColumns"]
        num = next(c["name"] for c in jcols if c["dtype"] in ("integer", "float"))
        totals = client.post(f"/queries/{joined}/totals", json={"columns": [num]}).json()
        preview = client.post(
            f"/workspaces/{ws}/queries/preview", json={"sourceId": deals, "definition": _joined_def(rel)}
        ).json()

    assert rows["dataVersions"] == {deals: 1}
    both = {deals: 1, accounts: 2}
    assert jrows["dataVersions"] == both
    assert agg["dataVersions"] == both
    assert totals["dataVersions"] == both
    assert preview["dataVersions"] == both
    validate_response("queries/rows-get.contract.yaml", 200, jrows)
    validate_response("queries/aggregate.contract.yaml", 200, agg)
    validate_response("queries/totals.contract.yaml", 200, totals)
    validate_response("queries/preview.contract.yaml", 200, preview)


@pytest.mark.unit
def test_the_query_export_carries_the_versions() -> None:
    with TestClient(app) as client:
        ws, deals, accounts = _seed(client)
        rel = _declare_id_join(client, ws, deals, accounts)
        joined = _create_join(client, ws, deals, rel).json()["id"]
        _refresh(client, ws, deals)
        xlsx = client.post(f"/queries/{joined}/export", json={"format": "xlsx"})
        csv = client.post(f"/queries/{joined}/export", json={"format": "csv"})

    both = {deals: 2, accounts: 1}
    assert json.loads(csv.headers["x-mdd-data-versions"]) == both
    assert json.loads(xlsx.headers["x-mdd-data-versions"]) == both
    cells = [[c.value for c in r] for r in load_workbook(io.BytesIO(xlsx.content))["Filters"].iter_rows()]
    i = next(k for k, r in enumerate(cells) if r[0] == "Data versions")
    assert cells[i + 1][:3] == ["dataset", "id", "version"]
    assert sorted(tuple(r[:3]) for r in cells[i + 2 : i + 4]) == sorted(
        [("deals", deals, 2), ("accounts", accounts, 1)]
    )


@pytest.mark.unit
def test_a_dataset_exports_its_visible_columns_and_its_version() -> None:
    with TestClient(app) as client:
        ws, deals, _accounts = _seed(client)
        client.patch(f"/datasets/{deals}/columns", json={"hidden": ["amount"]})
        _refresh(client, ws, deals)
        csv = client.post(f"/datasets/{deals}/export", json={"format": "csv"})
        xlsx = client.post(f"/datasets/{deals}/export", json={"format": "xlsx"})
        missing = client.post("/datasets/ds_00000000/export", json={"format": "csv"})
        bad = client.post(f"/datasets/{deals}/export", json={"format": "pdf"})

    assert csv.status_code == 200
    header = csv.content.decode("utf-8-sig").splitlines()[0].split(",")
    assert "amount" not in header and "id" in header
    assert json.loads(csv.headers["x-mdd-data-versions"]) == {deals: 2}
    assert csv.headers["content-disposition"] == 'attachment; filename="deals.csv"'
    wb = load_workbook(io.BytesIO(xlsx.content))
    meta = _meta(wb)
    assert meta["dataset"] == "deals" and meta["data_version"] == "2"
    assert missing.status_code == 404
    validate_response("datasets/export-post.contract.yaml", 404, missing.json())
    assert bad.status_code == 422


@pytest.mark.unit
def test_a_saved_copy_exports_as_frozen() -> None:
    with TestClient(app) as client:
        ws, deals, _accounts = _seed(client)
        _refresh(client, ws, deals)  # v2
        copy = client.post(f"/workspaces/{ws}/saved-copies", json={"label": "wk42", "datasetIds": [deals]}).json()[
            "datasets"
        ][0]
        _refresh(client, ws, deals)  # the original moves on to v3
        xlsx = client.post(f"/datasets/{copy['id']}/export", json={"format": "xlsx"})

    assert xlsx.status_code == 200
    assert json.loads(xlsx.headers["x-mdd-data-versions"]) == {copy["id"]: 2}
    meta = {
        r[0].value: r[1].value for r in load_workbook(io.BytesIO(xlsx.content))["Filters"].iter_rows() if r[0].value
    }
    assert meta["data_version"] == "2"
    assert "not the current numbers" in meta["frozen_copy_of"] and "deals" in meta["frozen_copy_of"]
    assert meta["label"] == "wk42"
