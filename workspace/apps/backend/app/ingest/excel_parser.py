"""Excel parse via openpyxl (read-only) + pandas dtype inference.

Two entry points:

- ``enumerate_sheets`` lists sheet names + cheap (row, column)
  counts + a sniffer-detected ``usedRange``. Used by ``POST /uploads``
  to power the wizard's Sheet-step UI without a full parse.
- ``parse_sheet`` reads one sheet — optionally a sub-range — and
  returns columns + row count + sample rows. Used by
  ``POST /uploads/{temp_id}/parse``.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Literal

import openpyxl
import pandas as pd

from app.ingest.csv_parser import ParseResult


SAMPLE_LIMIT = 10

Dtype = Literal["string", "integer", "float", "boolean", "date", "datetime"]


_RANGE_RE = re.compile(r"^([A-Z]+)([0-9]+):([A-Z]+)([0-9]+)$")


@dataclass
class SheetMeta:
    sheet: str
    row_count: int
    column_count: int
    used_range: str | None


def enumerate_sheets(path: Path) -> list[SheetMeta]:
    wb = openpyxl.load_workbook(filename=path, read_only=True, data_only=True)
    try:
        out: list[SheetMeta] = []
        for name in wb.sheetnames:
            ws = wb[name]
            dim = ws.calculate_dimension() if ws.max_row else None
            if dim and _RANGE_RE.match(dim):
                used = dim
            else:
                used = None
            row_count = max(0, (ws.max_row or 0) - 1) if used else 0
            col_count = ws.max_column or 0
            out.append(
                SheetMeta(
                    sheet=name,
                    row_count=row_count,
                    column_count=col_count,
                    used_range=used,
                )
            )
        return out
    finally:
        wb.close()


class ExcelParseError(Exception):
    pass


def _dtype_from_series(series: pd.Series) -> Dtype:
    kind = series.dtype.kind
    if kind == "b":
        return "boolean"
    if kind == "i" or kind == "u":
        return "integer"
    if kind == "f":
        return "float"
    if kind == "M":
        return "datetime"
    return "string"


_LEADING_ZERO = re.compile(r"^0\d+$")
_ZERO_SNIFF_ROWS = 200  # rows re-read as text to spot a leading zero


def _leading_zero_text_to_string(path: Path, kwargs: dict, df: pd.DataFrame, columns: list[dict]) -> None:
    """R175 — a column of digit TEXT with a leading zero (a phone `0968733440`, a
    code `00123`) is a STRING, not a number. pandas' Excel reader converts
    numeric-looking text to int64, so inference said `integer` and the commit
    wrote `968733440`; a join on the phone then matched almost nothing (dogfood:
    2 of 1,676 policies found their TNDS contract instead of 484). Re-read only the
    SAMPLE as text for the numeric-inferred columns and flip any column whose text
    carries a leading zero. Mutates ``columns`` in place."""
    numeric = [i for i, c in enumerate(columns) if c["dtype"] in ("integer", "float")]
    if not numeric:
        return
    sample_kwargs = {**kwargs, "dtype": str, "nrows": min(int(kwargs.get("nrows") or _ZERO_SNIFF_ROWS), _ZERO_SNIFF_ROWS)}
    try:
        text = pd.read_excel(path, **sample_kwargs)
    except Exception:  # pragma: no cover — the full read already succeeded
        return
    for i in numeric:
        if i >= text.shape[1]:
            continue
        cells = text.iloc[:, i].dropna().astype(str).str.strip()
        if cells.str.match(_LEADING_ZERO).any():
            columns[i]["dtype"] = "string"


def _parse_range(rng: str) -> tuple[str, str, int, int] | None:
    m = _RANGE_RE.match(rng)
    if not m:
        return None
    return m.group(1), m.group(3), int(m.group(2)), int(m.group(4))


def parse_sheet(
    path: Path,
    sheet: str,
    *,
    range_: str | None = None,
    has_header: bool = True,
) -> ParseResult:
    """Parse one sheet of an Excel workbook.

    Raises ``ExcelParseError`` for unknown-sheet / malformed-range /
    read failures. Caller maps to the contract's per-item failure
    shape.
    """

    try:
        kwargs: dict[str, object] = {
            "sheet_name": sheet,
            "header": 0 if has_header else None,
            "engine": "openpyxl",
        }
        if range_ is not None:
            parsed = _parse_range(range_)
            if parsed is None:
                raise ExcelParseError(f"range_invalid: {range_}")
            col_first, col_last, row_first, row_last = parsed
            kwargs["usecols"] = f"{col_first}:{col_last}"
            kwargs["skiprows"] = max(0, row_first - 1)
            kwargs["nrows"] = max(0, row_last - row_first + (0 if has_header else 1))
        df = pd.read_excel(path, **kwargs)
    except ExcelParseError:
        raise
    except ValueError as exc:
        raise ExcelParseError(f"parse_failed: {exc}") from exc
    except Exception as exc:  # pragma: no cover — defensive
        raise ExcelParseError(f"parse_failed: {exc}") from exc

    if not has_header:
        df.columns = [f"column{i + 1}" for i in range(len(df.columns))]
    else:
        df.columns = [str(c).strip() for c in df.columns]

    columns = [{"name": str(name), "dtype": _dtype_from_series(df[name])} for name in df.columns]
    _leading_zero_text_to_string(path, kwargs, df, columns)
    sample = df.head(SAMPLE_LIMIT)
    sample_rows: list[list[str | None]] = []
    for _, row in sample.iterrows():
        sample_rows.append([None if pd.isna(v) else str(v) for v in row.tolist()])
    return ParseResult(columns=columns, row_count=int(len(df)), sample_rows=sample_rows)
