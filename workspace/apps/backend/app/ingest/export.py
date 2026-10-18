"""R182 — export a table widget's numbers as CSV or XLSX.

The file is built over the SAME relation the table shows (the post-step,
dashboard-filtered result — `routers.queries._final_relation`), with every
row, not the visible page; the totals row (R181 spec) when the widget has
one; and, in XLSX, a second sheet naming the filters the numbers were
exported under. Cells keep their type: integers and floats are numbers, dates
are dates, a ratio column is a number formatted as a percent — Excel reads
them as such (purpose.md key design decision #3, Excel output first).

Caps at `EXPORT_MAX_ROWS` rows and says so (422 `export_too_large`) instead
of writing a file nobody can open.
"""

from __future__ import annotations

import csv
import io
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

EXPORT_MAX_ROWS = 100_000


def _csv_cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, datetime):
        return v.isoformat(sep=" ", timespec="seconds")
    if isinstance(v, date):
        return v.isoformat()
    return str(v)


def file_stem(title: str) -> str:
    """A file-name stem from a title (letters, digits, `-_`; spaces → `_`)."""
    return (
        "".join(ch if ch.isalnum() or ch in "-_ " else "_" for ch in title).strip().replace(" ", "_")[:60] or "export"
    )


def attachment_disposition(stem: str, ext: str) -> str:
    """`Content-Disposition` for a download. A non-Latin stem (e.g. Vietnamese)
    cannot ride in a plain header: an ASCII fallback plus RFC 5987 `filename*`.
    An ASCII stem gets exactly the plain form (R194)."""
    from urllib.parse import quote

    ascii_stem = stem.encode("ascii", "replace").decode("ascii").replace("?", "_")
    plain = f'attachment; filename="{ascii_stem}.{ext}"'
    if ascii_stem == stem:
        return plain
    return f"{plain}; filename*=UTF-8''{quote(f'{stem}.{ext}', safe='')}"


def build_csv(
    columns: list[dict],
    rows: list[list[Any]],
    totals: list[Any] | None,
    total_label: str,
    *,
    totalled: frozenset[str] = frozenset(),
) -> bytes:
    """UTF-8 with a BOM (so Excel opens it as UTF-8), header + rows + the
    totals row last when given (labelled — see `_label_totals`)."""
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow([c["name"] for c in columns])
    for r in rows:
        w.writerow([_csv_cell(v) for v in r])
    if totals is not None:
        w.writerow(_label_totals(columns, totals, total_label, totalled, as_text=True))
    return ("﻿" + buf.getvalue()).encode("utf-8")


def _label_totals(
    columns: list[dict], totals: list[Any], label: str, totalled: frozenset[str], *, as_text: bool
) -> list[Any]:
    """`Total` goes in the first column that has NO totals spec, so a file's
    totals row always says what it is (a CSV has no bold to tell it from a data
    row). A totalled column keeps its value even when it is empty (a ratio over
    0) — the label never overwrites a total (R191). Every column totalled → no
    label. (The page labels column 0 only; its footer row is styled, a file's is
    not.)"""
    out = list(totals)
    slot = next((i for i, c in enumerate(columns) if c["name"] not in totalled), None)
    if slot is not None:
        out[slot] = label
    return [(_csv_cell(v) if as_text else v) for v in out]


def build_xlsx(
    columns: list[dict],
    rows: list[list[Any]],
    totals: list[Any] | None,
    total_label: str,
    *,
    percent_columns: set[str],
    sheet_title: str,
    filters: list[tuple[str, list[str | None]]],
    meta: list[tuple[str, str]],
    filters_sheet_title: str,
    totalled: frozenset[str] = frozenset(),
    params: list[tuple[str, object, str]] | None = None,
    data_versions: list[tuple[str, str, int]] | None = None,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = _sheet_name(sheet_title)
    bold = Font(bold=True)
    ws.append([c["name"] for c in columns])
    for cell in ws[1]:
        cell.font = bold
    for r in rows:
        ws.append([_xlsx_cell(v) for v in r])
    if totals is not None:
        ws.append(_label_totals(columns, totals, total_label, totalled, as_text=False))
        for cell in ws[ws.max_row]:
            cell.font = bold
    # Number formats per column: percent for ratio columns, dates as dates.
    for ci, c in enumerate(columns, start=1):
        letter = get_column_letter(ci)
        fmt = None
        if c["name"] in percent_columns:
            fmt = "0.0%"
        elif c.get("dtype") == "date":
            fmt = "yyyy-mm-dd"
        elif c.get("dtype") == "datetime":
            fmt = "yyyy-mm-dd hh:mm:ss"
        if fmt:
            for cell in ws[letter][1:]:
                cell.number_format = fmt
        ws.column_dimensions[letter].width = max(10, min(40, len(c["name"]) + 2))
    ws.freeze_panes = "A2"

    fs = wb.create_sheet(_sheet_name(filters_sheet_title))
    for k, v in meta:
        fs.append([k, v])
    fs.append([])
    fs.append(["column", "values"])
    fs[fs.max_row][0].font = bold
    fs[fs.max_row][1].font = bold
    if filters:
        for col, values in filters:
            fs.append([col, ", ".join("(blank)" if v is None else str(v) for v in values)])
    else:
        fs.append(["—", "no filter"])
    # R194 (D7) — the parameter values the rows were read under, and where each
    # came from (viewer · dashboard · default). Only when the read resolved any.
    if params:
        fs.append([])
        fs.append(["Parameters"])
        fs[fs.max_row][0].font = bold
        fs.append(["name", "value", "source"])
        for cell in fs[fs.max_row]:
            cell.font = bold
        for name, value, source in params:
            fs.append([name, value, source])
        fs.column_dimensions["C"].width = 12
    # R196 (3.2c) — the data the rows were read from: each dataset's version.
    if data_versions:
        fs.append([])
        fs.append(["Data versions"])
        fs[fs.max_row][0].font = bold
        fs.append(["dataset", "id", "version"])
        for cell in fs[fs.max_row]:
            cell.font = bold
        for name, ds_id, version in data_versions:
            fs.append([name, ds_id, version])
        fs.column_dimensions["C"].width = 12
    fs.column_dimensions["A"].width = 24
    fs.column_dimensions["B"].width = 60

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _xlsx_cell(v: Any) -> Any:
    # openpyxl writes int / float / bool / date / datetime / str natively.
    return v


def _sheet_name(title: str) -> str:
    bad = set("[]:*?/\\")
    cleaned = "".join("_" if ch in bad else ch for ch in title).strip() or "Sheet"
    return cleaned[:31]
